"""Load model-linked locations into emtac.db from the QA location workbook."""

from __future__ import annotations

import argparse
import shutil
import sqlite3
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from build_emtac_hierarchy_from_bom import normalize_asset_number


DEFAULT_TARGET_DB = "data/emtac.db"
DEFAULT_WORKBOOK = (
    "outputs/qa_location_assets/"
    "xau_mech_locations_to_assets_setup_20260415_110055.xlsx"
)
DEFAULT_ORG_CODE = "XAU"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Load QA location values into emtac.db, linked many-to-one to model_id."
    )
    parser.add_argument("--target-db", default=DEFAULT_TARGET_DB)
    parser.add_argument("--workbook", default=DEFAULT_WORKBOOK)
    parser.add_argument("--org-code", default=DEFAULT_ORG_CODE)
    parser.add_argument(
        "--no-backup",
        action="store_true",
        help="Do not create a timestamped database backup before rebuilding location.",
    )
    return parser.parse_args()


def clean_text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def backup_database(db_path: Path) -> Path | None:
    if not db_path.exists():
        raise SystemExit(f"Target database does not exist: {db_path}")
    backup_path = db_path.with_suffix(
        f".before_locations_{datetime.now():%Y%m%d_%H%M%S}.db"
    )
    shutil.copy2(db_path, backup_path)
    return backup_path


def rebuild_location_table(connection: sqlite3.Connection) -> None:
    connection.execute("PRAGMA foreign_keys = OFF")
    connection.execute("DROP TABLE IF EXISTS location")
    connection.execute(
        """
        CREATE TABLE location (
            location_id INTEGER PRIMARY KEY,
            model_id INTEGER NOT NULL,
            location TEXT NOT NULL,
            location_description TEXT,
            UNIQUE (model_id, location),
            FOREIGN KEY (model_id)
                REFERENCES model (model_id)
        )
        """
    )
    connection.execute("PRAGMA foreign_keys = ON")


def load_asset_model_index(connection: sqlite3.Connection) -> dict[str, set[int]]:
    asset_columns = {
        row[1]
        for row in connection.execute("PRAGMA table_info(asset_number)").fetchall()
    }
    if "model_id" not in asset_columns:
        raise SystemExit("asset_number table is missing required model_id column")
    rows = connection.execute(
        """
        SELECT
            asset_number,
            model_id
        FROM asset_number
        WHERE model_id IS NOT NULL
        """
    ).fetchall()
    asset_to_model_ids: dict[str, set[int]] = defaultdict(set)
    for asset_number, model_id in rows:
        asset = normalize_asset_number(asset_number)
        if asset and model_id is not None:
            asset_to_model_ids[asset].add(int(model_id))
    return asset_to_model_ids


def workbook_rows(workbook_path: Path) -> tuple[list[str], list[dict[str, object]]]:
    if not workbook_path.exists():
        raise SystemExit(f"Workbook does not exist: {workbook_path}")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = [clean_text(cell.value) for cell in next(worksheet.iter_rows(max_row=1))]
    if any(header is None for header in headers):
        raise SystemExit("Workbook header row has blank column names")
    header_names = [str(header) for header in headers]
    rows = [
        dict(zip(header_names, row))
        for row in worksheet.iter_rows(min_row=2, values_only=True)
    ]
    workbook.close()
    return header_names, rows


def load_locations(
    target_db: Path,
    workbook_path: Path,
    org_code: str | None,
    make_backup: bool,
) -> dict[str, int | str | None]:
    backup_path = backup_database(target_db) if make_backup else None
    expected_org = org_code.strip().upper() if org_code else None

    headers, rows = workbook_rows(workbook_path)
    required = {"ORGANIZATION_CODE", "ASSET_NUMBER", "LOCATION_VALUE"}
    missing = required - set(headers)
    if missing:
        raise SystemExit(f"Workbook is missing column(s): {sorted(missing)}")

    connection = sqlite3.connect(target_db)
    try:
        connection.execute("PRAGMA foreign_keys = ON")
        connection.execute(
            "CREATE INDEX IF NOT EXISTS idx_asset_number_asset ON asset_number (asset_number)"
        )
        asset_to_model_ids = load_asset_model_index(connection)
        rebuild_location_table(connection)

        location_map: dict[tuple[int, str], str | None] = {}
        workbook_assets: set[str] = set()
        unmatched_assets: set[str] = set()
        multi_model_assets: set[str] = set()
        skipped_non_org = 0
        skipped_blank = 0

        for row in rows:
            if expected_org:
                row_org = (clean_text(row.get("ORGANIZATION_CODE")) or "").upper()
                if row_org != expected_org:
                    skipped_non_org += 1
                    continue

            asset = normalize_asset_number(clean_text(row.get("ASSET_NUMBER")))
            location = clean_text(row.get("LOCATION_VALUE"))
            if not asset or not location:
                skipped_blank += 1
                continue

            workbook_assets.add(asset)
            model_ids = asset_to_model_ids.get(asset)
            if not model_ids:
                unmatched_assets.add(asset)
                continue
            if len(model_ids) > 1:
                multi_model_assets.add(asset)

            location_description = clean_text(row.get("LOCATION_DESCRIPTION"))
            for model_id in model_ids:
                key = (model_id, location)
                existing_description = location_map.get(key)
                if existing_description is None and location_description is not None:
                    existing_description = location_description
                location_map[key] = existing_description

        inserts = [
            (model_id, location, description)
            for (model_id, location), description in sorted(location_map.items())
        ]
        connection.executemany(
            """
            INSERT INTO location (
                model_id,
                location,
                location_description
            )
            VALUES (?, ?, ?)
            """,
            inserts,
        )
        connection.commit()

        fk_issues = connection.execute("PRAGMA foreign_key_check").fetchall()
        return {
            "backup_path": str(backup_path) if backup_path else None,
            "workbook_rows": len(rows),
            "workbook_assets": len(workbook_assets),
            "skipped_non_org": skipped_non_org,
            "skipped_blank": skipped_blank,
            "unmatched_assets": len(unmatched_assets),
            "multi_model_assets": len(multi_model_assets),
            "location_rows": len(inserts),
            "model_ids_with_locations": len({model_id for model_id, _ in location_map}),
            "foreign_key_issues": len(fk_issues),
        }
    finally:
        connection.close()


def main() -> int:
    args = parse_args()
    stats = load_locations(
        Path(args.target_db),
        Path(args.workbook),
        args.org_code,
        not args.no_backup,
    )
    print(f"Target database: {args.target_db}")
    print(f"Workbook: {args.workbook}")
    if args.org_code:
        print(f"Organization filter: {args.org_code.strip().upper()}")
    for key, value in stats.items():
        print(f"{key}: {value}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
