"""Load an EMTAC Excel workbook back into emtac.db.

The loader expects workbook sheets that match the EMTAC table names:

- area
- equipment_group
- model
- asset_number
- location
- asset_part
- boms_for_department or boms_for_department_1, boms_for_department_2, ...

For hierarchy sheets, explicit primary-key columns are optional. If a sheet omits
its self ID column, the 1-based data-row number is used as the ID:

- area row 2 -> area_id 1
- equipment_group row 2 -> equipment_group_id 1
- model row 2 -> model_id 1

Foreign-key columns such as equipment_group.area_id, model.equipment_group_id,
asset_number.model_id, and location.model_id are still loaded from the workbook.
"""

from __future__ import annotations

import argparse
import re
import shutil
import sqlite3
from collections.abc import Iterable, Iterator
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


DEFAULT_DB = "data/emtac.db"
DEFAULT_WORKBOOK_GLOB = "outputs/emtac_export/emtac_db_export_*.xlsx"
LOAD_ORDER = [
    "area",
    "equipment_group",
    "model",
    "asset_number",
    "location",
    "asset_part",
    "boms_for_department",
]
DELETE_ORDER = [
    "asset_part",
    "location",
    "asset_number",
    "model",
    "equipment_group",
    "area",
    "boms_for_department",
]
PRIMARY_KEYS = {
    "area": "area_id",
    "equipment_group": "equipment_group_id",
    "model": "model_id",
    "asset_number": "asset_number_id",
    "location": "location_id",
    "asset_part": "asset_part_id",
}
BOM_SHEET_RE = re.compile(r"^boms_for_department(?:_\d+)?$", re.IGNORECASE)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Overwrite EMTAC SQLite tables from an Excel workbook export."
    )
    parser.add_argument(
        "--db",
        default=DEFAULT_DB,
        help="Target SQLite database. Default: data/emtac.db",
    )
    parser.add_argument(
        "--workbook",
        default=None,
        help=(
            "Excel workbook to load. If omitted, the newest "
            f"{DEFAULT_WORKBOOK_GLOB} file is used."
        ),
    )
    parser.add_argument(
        "--no-backup",
        action="store_true",
        help="Do not create a timestamped database backup before loading.",
    )
    parser.add_argument(
        "--skip-boms",
        action="store_true",
        help="Do not load boms_for_department sheets.",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=5000,
        help="Rows per SQLite executemany batch.",
    )
    return parser.parse_args()


def clean_header(value: object) -> str:
    return str(value).strip() if value is not None else ""


def clean_value(value: object) -> object:
    if isinstance(value, str):
        text = value.strip()
        return text if text else None
    return value


def backup_database(db_path: Path) -> Path:
    if not db_path.exists():
        raise SystemExit(f"Database does not exist: {db_path}")
    backup_path = db_path.with_name(
        f"{db_path.stem}_before_excel_load_{datetime.now():%Y%m%d_%H%M%S}{db_path.suffix}"
    )
    shutil.copy2(db_path, backup_path)
    return backup_path


def resolve_workbook(path_text: str | None) -> Path:
    if path_text:
        return Path(path_text)
    candidates = sorted(
        Path().glob(DEFAULT_WORKBOOK_GLOB),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise SystemExit(
            "No workbook was supplied and no export workbook was found matching "
            f"{DEFAULT_WORKBOOK_GLOB}"
        )
    return candidates[0]


def table_columns(connection: sqlite3.Connection, table_name: str) -> list[str]:
    rows = connection.execute(f'PRAGMA table_info("{table_name}")').fetchall()
    if not rows:
        raise SystemExit(f"Table does not exist in target database: {table_name}")
    return [str(row[1]) for row in rows]


def worksheet_headers(worksheet) -> list[str]:
    try:
        header_row = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        raise SystemExit(f"Worksheet is empty: {worksheet.title}") from None
    headers = [clean_header(value) for value in header_row]
    if not any(headers):
        raise SystemExit(f"Worksheet has no headers: {worksheet.title}")
    return headers


def workbook_table_sheets(workbook, table_name: str) -> list[str]:
    if table_name == "boms_for_department":
        return sorted(
            [name for name in workbook.sheetnames if BOM_SHEET_RE.match(name)],
            key=lambda name: (
                0
                if name == "boms_for_department"
                else int(name.rsplit("_", 1)[1])
                if name.rsplit("_", 1)[-1].isdigit()
                else 9999,
                name,
            ),
        )
    return [table_name] if table_name in workbook.sheetnames else []


def normalized_rows(
    worksheet,
    table_name: str,
    db_columns: list[str],
) -> Iterator[tuple[object, ...]]:
    headers = worksheet_headers(worksheet)
    header_positions = {header: idx for idx, header in enumerate(headers) if header}
    pk_column = PRIMARY_KEYS.get(table_name)

    missing = [
        column
        for column in db_columns
        if column not in header_positions and column != pk_column
    ]
    if missing:
        raise SystemExit(
            f"Worksheet {worksheet.title} is missing required column(s) for "
            f"{table_name}: {missing}"
        )

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True),
        start=1,
    ):
        if row_is_blank(row):
            continue

        values: list[object] = []
        for column in db_columns:
            if column in header_positions:
                value = row[header_positions[column]]
            elif column == pk_column:
                value = row_number
            else:
                value = None
            values.append(clean_value(value))
        yield tuple(values)


def row_is_blank(row: Iterable[object]) -> bool:
    return all(value is None or (isinstance(value, str) and not value.strip()) for value in row)


def batched(rows: Iterable[tuple[object, ...]], batch_size: int) -> Iterator[list[tuple[object, ...]]]:
    batch: list[tuple[object, ...]] = []
    for row in rows:
        batch.append(row)
        if len(batch) >= batch_size:
            yield batch
            batch = []
    if batch:
        yield batch


def insert_rows(
    connection: sqlite3.Connection,
    table_name: str,
    columns: list[str],
    rows: Iterable[tuple[object, ...]],
    batch_size: int,
) -> int:
    placeholders = ", ".join("?" for _ in columns)
    column_sql = ", ".join(f'"{column}"' for column in columns)
    sql = f'INSERT INTO "{table_name}" ({column_sql}) VALUES ({placeholders})'
    count = 0
    for batch in batched(rows, batch_size):
        connection.executemany(sql, batch)
        count += len(batch)
    return count


def reset_sqlite_sequence(
    connection: sqlite3.Connection,
    table_name: str,
    pk_column: str | None,
) -> None:
    if not pk_column:
        return
    if not connection.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='sqlite_sequence'"
    ).fetchone():
        return
    max_id = connection.execute(
        f'SELECT COALESCE(MAX("{pk_column}"), 0) FROM "{table_name}"'
    ).fetchone()[0]
    connection.execute("DELETE FROM sqlite_sequence WHERE name = ?", (table_name,))
    connection.execute(
        "INSERT INTO sqlite_sequence (name, seq) VALUES (?, ?)",
        (table_name, max_id),
    )


def load_workbook_to_db(
    db_path: Path,
    workbook_path: Path,
    make_backup: bool,
    skip_boms: bool,
    batch_size: int,
) -> dict[str, object]:
    if not workbook_path.exists():
        raise SystemExit(f"Workbook does not exist: {workbook_path}")
    backup_path = backup_database(db_path) if make_backup else None

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    connection = sqlite3.connect(db_path)
    loaded_counts: dict[str, int] = {}
    try:
        available_sheets = {
            table_name: workbook_table_sheets(workbook, table_name)
            for table_name in LOAD_ORDER
        }
        if skip_boms:
            available_sheets["boms_for_department"] = []

        core_tables = [
            "area",
            "equipment_group",
            "model",
            "asset_number",
            "location",
        ]
        present_core = [
            table_name for table_name in core_tables if available_sheets[table_name]
        ]
        if present_core and len(present_core) != len(core_tables):
            missing_core = [
                table_name
                for table_name in core_tables
                if not available_sheets[table_name]
            ]
            raise SystemExit(
                "Workbook has a partial EMTAC hierarchy. To overwrite hierarchy "
                f"tables it must include all core sheets. Missing: {missing_core}"
            )

        tables_to_delete: list[str] = []
        if present_core:
            tables_to_delete.extend(
                [
                    "asset_part",
                    "location",
                    "asset_number",
                    "model",
                    "equipment_group",
                    "area",
                ]
            )
        if available_sheets["boms_for_department"]:
            tables_to_delete.append("boms_for_department")

        connection.execute("PRAGMA foreign_keys = OFF")
        connection.execute("BEGIN")

        for table_name in DELETE_ORDER:
            if table_name not in tables_to_delete:
                continue
            connection.execute(f'DELETE FROM "{table_name}"')

        for table_name in LOAD_ORDER:
            sheet_names = available_sheets[table_name]
            if not sheet_names:
                loaded_counts[table_name] = 0
                continue

            columns = table_columns(connection, table_name)
            total = 0
            for sheet_name in sheet_names:
                worksheet = workbook[sheet_name]
                rows = normalized_rows(worksheet, table_name, columns)
                total += insert_rows(connection, table_name, columns, rows, batch_size)
            loaded_counts[table_name] = total
            reset_sqlite_sequence(
                connection,
                table_name,
                PRIMARY_KEYS.get(table_name),
            )

        connection.commit()
        connection.execute("PRAGMA foreign_keys = ON")
        fk_issues = connection.execute("PRAGMA foreign_key_check").fetchall()
    except Exception:
        connection.rollback()
        raise
    finally:
        workbook.close()
        connection.close()

    return {
        "backup_path": str(backup_path) if backup_path else None,
        "loaded_counts": loaded_counts,
        "foreign_key_issues": len(fk_issues),
    }


def main() -> int:
    args = parse_args()
    workbook_path = resolve_workbook(args.workbook)
    result = load_workbook_to_db(
        Path(args.db),
        workbook_path,
        not args.no_backup,
        args.skip_boms,
        args.batch_size,
    )
    print(f"Database: {args.db}")
    print(f"Workbook: {workbook_path}")
    if result["backup_path"]:
        print(f"Backup: {result['backup_path']}")
    for table_name, count in result["loaded_counts"].items():
        print(f"{table_name}: {count}")
    print(f"foreign_key_issues: {result['foreign_key_issues']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
