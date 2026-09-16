"""Build auditable asset, area and combined work-order summaries; no assumed MTBF."""
import csv
from collections import defaultdict
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/lifecare_mtbf_12_months'
START, END = '2025-09-09', '2026-09-09'

def weekday_hours(start, end):
    """Count only Mon-Fri portions of a half-open interval, on Oracle's clock."""
    if end < start:
        raise ValueError('Interval end precedes start')
    total = 0.0
    cursor = start
    while cursor < end:
        boundary = min(end, cursor.replace(hour=0, minute=0, second=0, microsecond=0) + timedelta(days=1))
        if cursor.weekday() < 5:
            total += (boundary-cursor).total_seconds()/3600
        cursor = boundary
    return total

def read(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def summarize():
    assets = read(ROOT / 'outputs/lifecare_mtbf_assets/lifecare_asset_review_20260909.csv')
    raw = read(OUT / 'work_orders.csv')
    numbers = [a['Asset Number'] for a in assets]
    assert len(numbers) == len(set(numbers)), 'Resolve cross-area asset assignments before rollup'
    by_asset = defaultdict(dict)
    for r in raw:
        by_asset[r['ASSET_NUMBER']].setdefault(r['WIP_ENTITY_ID'], []).append(r)

    def counts(groups):
        created, cancelled, pm, dm, other, coded = (set() for _ in range(6))
        for wo, records in groups.items():
            r = records[0]
            in_period = START <= r['CREATED_DATE'] < END
            cancel = 'CANCEL' in r['WORK_ORDER_STATUS'].upper()
            if in_period:
                created.add(wo)
                if cancel:
                    cancelled.add(wo)
                elif r['WORK_ORDER'].upper().startswith('PM') or r['WO_TYPE'].upper() in ('PM', 'PREVENTIVE', 'PREVENTIVE MAINTENANCE'):
                    pm.add(wo)
                elif r['WO_TYPE'].upper() == 'DM':
                    dm.add(wo)
                else:
                    other.add(wo)
            if not cancel:
                coded.update(x['FAILURE_ID'] for x in records if x['FAILURE_ID'] and START <= x['FAILURE_DATE'] < END)
        assert len(created) == sum(map(len, (cancelled, pm, dm, other)))
        return [len(x) for x in (created, cancelled, pm, dm, other, coded)]

    def timing(groups):
        dates = sorted(datetime.fromisoformat(records[0]['CREATED_DATE'])
                       for wo, records in groups.items() if counts({wo: records})[3])
        weekend_count = sum(d.weekday() >= 5 for d in dates)
        if len(dates) < 2:
            return [0, None, dates[0].isoformat(' ') if dates else None,
                    dates[-1].isoformat(' ') if dates else None,
                    'No DM work orders' if not dates else 'Only one DM; no interval',
                    None, 0, weekend_count, len(dates)-weekend_count]
        gaps = [(b-a).total_seconds()/3600 for a,b in zip(dates, dates[1:])]
        mean = sum(gaps)/len(gaps)
        assert abs(mean - (dates[-1]-dates[0]).total_seconds()/3600/(len(dates)-1)) < 1e-8
        weekdays = weekday_hours(dates[0], dates[-1])
        return [len(gaps), mean, dates[0].isoformat(' '), dates[-1].isoformat(' '),
                'Mon-Fri timing proxy; weekend DMs retained', weekdays/len(gaps),
                sum(gaps)-weekdays, weekend_count, len(dates)-weekend_count]

    headers = ['Area','Asset Number','Asset Description','Total WOs Created','Cancelled WOs','PM WOs (non-cancelled)','DM WOs (non-cancelled)','Other WOs (non-cancelled)','Recorded Failure IDs in Period','Failure Count (agreed DM basis)','Observed DM Intervals','Mean Calendar Hours Between DMs','First Included DM Created','Last Included DM Created','Timing Status','Mean Weekday Hours Between DMs','Weekend Hours Removed from Observed Span','Weekend DMs (review)','Weekday DMs','Review Note']
    asset_rows=[]
    for a in assets:
        totals = counts(by_asset[a['Asset Number']])
        asset_rows.append([a['Area'],a['Asset Number'],a['Asset Description'],*totals,totals[3],*timing(by_asset[a['Asset Number']]),a['Review Note']])
    summary_headers=['Scope','Assets in Scope',*headers[3:-1]]
    summaries=[]
    for area in ['Bagfab','Filling','Overwrap','Packout','Combined']:
        selected=[a for a in assets if area == 'Combined' or a['Area']==area]
        groups={k:v for a in selected for k,v in by_asset[a['Asset Number']].items()}
        totals = counts(groups)
        summaries.append([area,len(selected),*totals,totals[3],*timing(groups)])
    for i in range(2,9):
        assert sum(r[i] for r in summaries[:4]) == summaries[4][i], 'Rollup mismatch'
    w=Workbook(); notes=w.active;notes.title='Read Me'
    for note in [
        'Lifecare: last 12 months by asset, by area, and combined',
        'Period: 2025-09-09 00:00 inclusive to 2026-09-09 00:00 exclusive; Oracle date basis.',
        'Live Oracle work-order extraction; scope is the 683-record preliminary April asset list.',
        'Includes all scoped assets, including assets with zero work orders. Current active status is not verified.',
        'Counts use distinct WIP_ENTITY_ID; multiple failure-code rows do not multiply work orders.',
        'WO counts use creation date; failure-ID counts use failure date and exclude cancelled work orders.',
        'PM: PM-prefixed WO or PM/Preventive/Preventive Maintenance type. DM: exact DM type, excluding PM and cancelled.',
        'User-approved failure basis: each distinct non-cancelled DM work order counts as one failure; PMs are excluded.',
        'Other work-order types are excluded from the agreed DM failure count. All WO categories remain visible for reconciliation.',
        'DM-based failure counts are a reporting proxy, not independently verified breakdown events.',
        'User selected Date Created as the timing basis for comparisons. The displayed measure is mean calendar hours between DM work orders.',
        'Formula: (last included creation timestamp - first included creation timestamp) in hours / (included DM count - 1).',
        'Weekend-adjusted formula: Mon-Fri hours between first and last included DM creation / (included DM count - 1).',
        'Saturday 00:00 through Monday 00:00 is excluded from elapsed hours, including partial weekend endpoints.',
        'Mon-Fri includes all 24 hours per day. Holidays, weekday shutdowns and shift breaks have not been subtracted.',
        'Weekend-created DMs remain included under the agreed non-cancelled DM rule and are flagged for review; creation on a weekend does not prove weekend production.',
        'Two DMs entirely within the same weekend can have a zero weekday interval. This is not a claim of zero actual runtime.',
        'At least two included DMs are needed. Zero or one yields a blank mean with an explanatory status.',
        'Only consecutive intervals within the reporting window are used; leading/trailing gaps to window boundaries are not included.',
        'Distinct work orders with identical timestamps remain separate events and contribute zero-length intervals.',
        'Maintenance labor hours and work-order elapsed time are not machine operating hours or measured downtime.',
        'Area and combined means use consecutive DM creation timestamps across all assets in that scope; they are not averages of asset means.',
        'Area and combined measures describe DM arrival frequency; scopes with more assets can have shorter intervals. Compare the same scope over time.',
        'This is a calendar-based MTBF proxy, not operating-hour MTBF or measured time between area-stopping failures. Actual failure timing may differ from creation time.',
        'Parent/child assets, controls, software and expired records require boundary review before MTBF aggregation.',
    ]: notes.append([note])
    notes.column_dimensions['A'].width=145
    def sheet(name, head, rows):
        s=w.create_sheet(name);s.append(head)
        for row in rows:
            s.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in row])
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
        for c in s[1]: c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='24476A')
        for cells in s.columns:
            s.column_dimensions[cells[0].column_letter].width=min(65,max(18,len(str(cells[0].value))+3))
    sheet('By Asset',headers,asset_rows)
    for area in ['Bagfab','Filling','Overwrap','Packout']:
        sheet(area+' Assets', headers, [r for r in asset_rows if r[0] == area])
    sheet('By Area',summary_headers,summaries[:4])
    sheet('Combined',summary_headers,summaries[4:])
    area_map={a['Asset Number']:a['Area'] for a in assets}
    sheet('Work Order Detail',['Area',*raw[0].keys()] if raw else ['Area'],[[area_map[r['ASSET_NUMBER']],*r.values()] for r in raw])
    included = {}
    for r in raw:
        if counts({r['WIP_ENTITY_ID']: [r]})[3]:
            included.setdefault(r['WIP_ENTITY_ID'], r)
    assert len(included) == summaries[4][8]
    sheet('Included DM Work Orders',['Area',*raw[0].keys()] if raw else ['Area'],[[area_map[r['ASSET_NUMBER']],*r.values()] for r in included.values()])
    weekend = [r for r in included.values() if datetime.fromisoformat(r['CREATED_DATE']).weekday() >= 5]
    sheet('Weekend DM Review', ['Area',*raw[0].keys()] if raw else ['Area'], [[area_map[r['ASSET_NUMBER']],*r.values()] for r in weekend])
    interval_rows = []
    for a in assets:
        events = sorted([r for r in included.values() if r['ASSET_NUMBER'] == a['Asset Number']], key=lambda r:(r['CREATED_DATE'],r['WIP_ENTITY_ID']))
        asset_gaps=[]
        for previous, current in zip(events,events[1:]):
            start, end = (datetime.fromisoformat(r['CREATED_DATE']) for r in (previous,current))
            calendar = (end-start).total_seconds()/3600
            weekday = weekday_hours(start,end)
            assert 0 <= weekday <= calendar + 1e-8
            asset_gaps.append(weekday)
            interval_rows.append([a['Area'],a['Asset Number'],previous['WORK_ORDER'],current['WORK_ORDER'],previous['CREATED_DATE'],current['CREATED_DATE'],calendar,calendar-weekday,weekday,'Yes' if end.weekday() >= 5 else 'No'])
        if asset_gaps:
            assert abs(sum(asset_gaps)-weekday_hours(datetime.fromisoformat(events[0]['CREATED_DATE']),datetime.fromisoformat(events[-1]['CREATED_DATE']))) < 1e-7
    sheet('Asset DM Intervals',['Area','Asset Number','Previous DM','Current DM','Previous Created','Current Created','Calendar Hours','Weekend Hours Removed','Weekday Hours','Current DM Created on Weekend'],interval_rows)
    assert len(interval_rows) == sum(r[10] for r in asset_rows)
    assert len(weekend) == sum(r[17] for r in asset_rows)
    for name,head,rows in [('by_asset',headers,asset_rows),('by_area',summary_headers,summaries[:4]),('combined',summary_headers,summaries[4:])]:
        with (OUT/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.writer(f);writer.writerow(head);writer.writerows(rows)
    path=OUT/'lifecare_12_month_review_20260909.xlsx';w.save(path)
    check=load_workbook(path,read_only=True)
    assert check['By Asset'].max_row-1 == len(assets)
    print(summary_headers)
    for r in summaries: print(r)
    print(path)

if __name__ == '__main__':
    summarize()
