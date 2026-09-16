"""Package the validated DM review into main, area and individual asset sheets."""
import csv
import re
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from summarize_lifecare_12_month_review import weekday_hours

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'outputs/lifecare_mtbf_12_months'
AREAS = ['Bagfab', 'Filling', 'Overwrap', 'Packout']

def read(name):
    with (OUT / name).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def number(value):
    return float(value) if value else None

def safe(value):
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value

def build():
    assets = read('by_asset.csv')
    area_rows = read('by_area.csv') + read('combined.csv')
    events = defaultdict(dict)
    for r in read('work_orders.csv'):
        if ('2025-09-09' <= r['CREATED_DATE'] < '2026-09-09'
            and r['WO_TYPE'].upper() == 'DM'
            and not r['WORK_ORDER'].upper().startswith('PM')
            and 'CANCEL' not in r['WORK_ORDER_STATUS'].upper()):
            events[r['ASSET_NUMBER']].setdefault(r['WIP_ENTITY_ID'], r)

    w = Workbook()
    main = w.active
    main.title = 'Main Summary'
    for area in AREAS:
        w.create_sheet(area)
    names = {}
    used = {s.lower() for s in w.sheetnames}
    for a in assets:
        base = re.sub(r'[\\/*?:\[\]]', '_', a['Asset Number']).strip("'") or 'Asset'
        title = base[:31]
        suffix = 1
        while title.lower() in used:
            ending = f'_{suffix}'
            title = base[:31-len(ending)] + ending
            suffix += 1
        used.add(title.lower())
        names[a['Asset Number']] = title
        w.create_sheet(title)

    links = []
    def link(s, address, target, label):
        c = s[address]
        c.value = safe(label)
        c.hyperlink = "#'" + target.replace("'", "''") + "'!A1"
        c.font = Font(color='0563C1', underline='single')
        links.append(target)

    def title(s, text):
        s['A1'] = safe(text)
        s['A1'].font = Font(size=18, bold=True, color='24476A')
        s.sheet_view.showGridLines = False
        s.sheet_properties.pageSetUpPr.fitToPage = True
        s.page_setup.orientation = 'landscape'
        s.page_setup.paperSize = s.PAPERSIZE_A3
        s.page_setup.fitToWidth = 1
        s.page_setup.fitToHeight = 0

    def table(s, row, headers, records, widths):
        for col, h in enumerate(headers, 1):
            c = s.cell(row, col, h)
            c.fill = PatternFill('solid', fgColor='24476A')
            c.font = Font(bold=True, color='FFFFFF')
            c.alignment = Alignment(wrap_text=True, vertical='center')
        s.row_dimensions[row].height = 36
        for i, record in enumerate(records, row+1):
            for col, value in enumerate(record, 1):
                c = s.cell(i, col, safe(value))
                c.alignment = Alignment(vertical='top', wrap_text=col in (2, 3))
                if isinstance(value, float):
                    c.number_format = '0.00'
                if i % 2 == 0:
                    c.fill = PatternFill('solid', fgColor='EFF4F8')
        for col, width in enumerate(widths, 1):
            s.column_dimensions[get_column_letter(col)].width = width
        s.auto_filter.ref = f'A{row}:{get_column_letter(len(headers))}{max(row,len(records)+row)}'
        s.freeze_panes = f'C{row+1}'
        s.print_title_rows = f'1:{row}'

    title(main, 'Lifecare — 12-month DM comparison')
    main['A3'] = 'September 9, 2025–September 8, 2026 | Date Created basis | Saturdays and Sundays excluded from elapsed hours'
    overview = [[r['Scope'],int(r['Assets in Scope']),int(r['Failure Count (agreed DM basis)']),
                 number(r['Mean Weekday Hours Between DMs']),number(r['Mean Calendar Hours Between DMs']),
                 int(r['Weekend DMs (review)']),int(r['Observed DM Intervals'])] for r in area_rows]
    table(main, 5, ['Area / Scope','Asset Records','Included DMs','Mean Weekday Hours Between DMs','Mean Calendar Hours Between DMs','Weekend DMs to Review','Observed Intervals'], overview, [24,20,20,28,28,24,24])
    for row, area in enumerate(AREAS,6):
        link(main,f'A{row}',area,area+' — open assets')
    notes = [
        'Navigation: select an area above, then select an asset number on its area sheet. Each asset sheet has links back.',
        'Included events: distinct non-cancelled DM work orders. PMs and other work-order types are excluded.',
        'Mean weekday hours = Mon–Fri hours between first and last included DM creation / (DM count − 1).',
        'All 24 weekday hours count. Holidays, weekday shutdowns and shift breaks are not removed.',
        'Weekend DMs remain included and are flagged. Weekend-created records do not establish that production was running.',
        'Zero or one included DM: mean interval is blank because fewer than two events were observed.',
        'Area and combined results measure consecutive DM arrivals across the scope, not the average of asset means.',
        'This is a schedule-adjusted comparison proxy, not measured operating-hour MTBF. Keep asset scope consistent when comparing periods.',
        'Asset scope: 683 records from the preliminary April 2026 list, including controls/software and records needing active-status review.',
        'Source: live Oracle work orders extracted September 9, 2026. Observed intervals omit leading/trailing window gaps.',
    ]
    for row, note in enumerate(notes,14):
        main.merge_cells(start_row=row,start_column=1,end_row=row,end_column=7)
        c=main.cell(row,1,note)
        c.alignment=Alignment(wrap_text=True,vertical='top')
        main.row_dimensions[row].height=32

    asset_headers = ['Asset Number','Asset Description','Included DMs','Mean Weekday Hours Between DMs','Mean Calendar Hours Between DMs','Weekend DMs','Observed Intervals','Weekend Hours Removed','First DM Created','Last DM Created','Review Note']
    for area in AREAS:
        s=w[area]
        title(s,area+' — individual assets')
        link(s,'A2','Main Summary','Back to Main Summary')
        selected=[a for a in assets if a['Area']==area]
        records=[[a['Asset Number'],a['Asset Description'],int(a['Failure Count (agreed DM basis)']),
                  number(a['Mean Weekday Hours Between DMs']),number(a['Mean Calendar Hours Between DMs']),
                  int(a['Weekend DMs (review)']),int(a['Observed DM Intervals']),
                  number(a['Weekend Hours Removed from Observed Span']),a['First Included DM Created'],
                  a['Last Included DM Created'],a['Review Note']] for a in selected]
        table(s,4,asset_headers,records,[27,65,16,27,27,18,20,24,23,23,65])
        for row,a in enumerate(selected,5):
            link(s,f'A{row}',names[a['Asset Number']],a['Asset Number'])

    detail_count=0
    for a in assets:
        s=w[names[a['Asset Number']]]
        title(s,a['Asset Number'])
        link(s,'A2','Main Summary','Back to Main Summary')
        link(s,'C2',a['Area'],'Back to '+a['Area'])
        s['A3']='Description'
        s['B3']=safe(a['Asset Description'])
        for row,(label,value) in enumerate([
            ('Area',a['Area']),('Included non-cancelled DMs',int(a['Failure Count (agreed DM basis)'])),
            ('Mean weekday hours between DMs',number(a['Mean Weekday Hours Between DMs'])),
            ('Mean calendar hours between DMs',number(a['Mean Calendar Hours Between DMs'])),
            ('Weekend-created DMs to review',int(a['Weekend DMs (review)'])),
            ('Observed intervals',int(a['Observed DM Intervals'])),
            ('Weekend hours removed',number(a['Weekend Hours Removed from Observed Span'])),
            ('Timing status',a['Timing Status']),('Asset review note',a['Review Note']),
        ],4):
            s.cell(row,1,label)
            c=s.cell(row,2,safe(value))
            if isinstance(value,float): c.number_format='0.00'
        s['A14']='Below: each included DM, with elapsed hours since the preceding included DM on this asset.'
        ordered=sorted(events[a['Asset Number']].values(),key=lambda r:(r['CREATED_DATE'],r['WIP_ENTITY_ID']))
        assert len(ordered)==int(a['Failure Count (agreed DM basis)'])
        records=[]
        previous=None
        weekday_gaps=[]
        for r in ordered:
            now=datetime.fromisoformat(r['CREATED_DATE'])
            calendar=(now-previous).total_seconds()/3600 if previous is not None else None
            weekday=weekday_hours(previous,now) if previous is not None else None
            if weekday is not None: weekday_gaps.append(weekday)
            records.append([r['WORK_ORDER'],r['CREATED_DATE'],r['WO_DESCRIPTION'],r['WORK_ORDER_STATUS'],now.strftime('%A'),
                            'Yes — review' if now.weekday()>=5 else 'No',calendar,
                            calendar-weekday if calendar is not None else None,weekday])
            previous=now
        if weekday_gaps:
            assert abs(sum(weekday_gaps)/len(weekday_gaps)-float(a['Mean Weekday Hours Between DMs']))<1e-7
        detail_count+=len(records)
        table(s,16,['DM Work Order','Date Created','Work Description','Status','Day Created','Created on Weekend','Calendar Hours Since Prior DM','Weekend Hours Removed','Weekday Hours Since Prior DM'],records,[40,65,75,22,18,23,27,25,27])
        s.freeze_panes='C17'
        s.sheet_properties.tabColor={'Bagfab':'548235','Filling':'4472C4','Overwrap':'BF9000','Packout':'7030A0'}[a['Area']]

    assert detail_count==int(area_rows[-1]['Failure Count (agreed DM basis)'])
    assert all(target in w.sheetnames for target in links)
    assert len(w.sheetnames)==5+len(assets)
    path=OUT/'Lifecare_MTBF_Main_Areas_Each_Asset.xlsx'
    w.save(path)
    check=load_workbook(path,read_only=True)
    assert check.sheetnames[:5]==['Main Summary',*AREAS]
    assert len(check.sheetnames)==688
    assert sum(max(0,check[names[a['Asset Number']]].max_row-16) for a in assets)==detail_count
    check.close()
    print(f'Saved {path}\nVerified {len(assets)} asset sheets, 4 area sheets, 1 main sheet, {detail_count} DM detail rows and {len(links)} navigation links.')

if __name__=='__main__':
    build()
